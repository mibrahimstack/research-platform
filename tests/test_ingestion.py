from docx import Document

from ingestion.batch_ingest import discover_documents, ingest_documents
from ingestion.parse_paper import parse_paper_file, parse_paper_xml, paper_to_chunks


def test_parse_paper_xml_returns_title_and_sections(fixtures_dir):
    parsed = parse_paper_xml(fixtures_dir / "sample_paper.xml")

    assert parsed["title"] == "Sample Paper"
    assert "new treatment for diabetes" in parsed["abstract"]
    assert parsed["sections"][0]["heading"] == "Introduction"
    assert "improve outcomes" in parsed["sections"][0]["text"]


def test_paper_to_chunks_creates_expected_chunks(fixtures_dir):
    parsed = parse_paper_xml(fixtures_dir / "sample_paper.xml")
    chunks = paper_to_chunks(parsed)

    assert len(chunks) >= 2
    assert any(chunk["section"] == "Abstract" for chunk in chunks)
    assert any(chunk["section"] == "Introduction" for chunk in chunks)
    assert all("text" in chunk for chunk in chunks)


def test_parse_paper_file_handles_docx(tmp_path):
    doc_path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Sample DOCX Paper", level=1)
    document.add_paragraph("This document discusses a diabetes treatment study.")
    document.save(doc_path)

    parsed = parse_paper_file(doc_path)

    assert parsed["title"] == "Sample DOCX Paper"
    assert parsed["abstract"] or parsed["sections"]
    assert any("diabetes treatment" in section["text"] for section in parsed["sections"])


def test_batch_ingestion_writes_manifest_and_failure_report(tmp_path, fixtures_dir):
    papers_dir = tmp_path / "papers"
    output_dir = tmp_path / "processed"
    papers_dir.mkdir()
    (papers_dir / "valid.xml").write_text((fixtures_dir / "sample_paper.xml").read_text(encoding="utf-8"), encoding="utf-8")
    (papers_dir / "broken.pdf").write_bytes(b"not a PDF")

    expected_chunks = len(paper_to_chunks(parse_paper_xml(papers_dir / "valid.xml")))
    summary = ingest_documents(papers_dir, output_dir)

    assert summary == {"documents": 2, "chunks": expected_chunks, "failed": 1}
    assert (output_dir / "chunks.jsonl").exists()
    assert (output_dir / "ingestion_manifest.jsonl").exists()
    assert (output_dir / "ingestion_failures.jsonl").exists()
